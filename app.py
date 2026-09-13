from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
import sqlite3
import os
import functools
from io import BytesIO
from datetime import datetime
from werkzeug.security import generate_password_hash, check_password_hash

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# PostgreSQL
try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
except ImportError:
    psycopg2 = None
    RealDictCursor = None


app = Flask(__name__)

# Needed to keep login sessions secure.
# On Render, set a SECRET_KEY environment variable
# so sessions survive restarts/deploys.
app.secret_key = os.environ.get(
    "SECRET_KEY",
    "studypoint-dev-secret-change-this"
)


# =====================================================
# DATABASE SETTINGS
# =====================================================

DATABASE = "database.db"

# Render PostgreSQL provides DATABASE_URL
DATABASE_URL = os.environ.get("DATABASE_URL")

# True = PostgreSQL
# False = SQLite
USE_POSTGRES = bool(DATABASE_URL)

print("DATABASE_URL exists:", bool(DATABASE_URL))
print("USE_POSTGRES:", USE_POSTGRES)


# =====================================================
# DATABASE CONNECTION WRAPPER
# =====================================================

class DatabaseConnection:

    def __init__(self):

        self.is_postgres = USE_POSTGRES

        if self.is_postgres:

            if psycopg2 is None:
                raise RuntimeError(
                    "psycopg2-binary is not installed."
                )

            # Some services may provide postgres://
            # psycopg2 works better with postgresql://
            database_url = DATABASE_URL

            if database_url.startswith("postgres://"):
                database_url = database_url.replace(
                    "postgres://",
                    "postgresql://",
                    1
                )

            self.conn = psycopg2.connect(
                database_url,
                cursor_factory=RealDictCursor
            )

        else:

            self.conn = sqlite3.connect(DATABASE)

            self.conn.row_factory = sqlite3.Row


    def execute(self, query, params=()):

        # PostgreSQL uses %s
        if self.is_postgres:

            query = query.replace("?", "%s")

            cursor = self.conn.cursor()
            cursor.execute(query, params)

            return cursor

        # SQLite uses ?
        return self.conn.execute(query, params)


    def commit(self):

        self.conn.commit()


    def close(self):

        self.conn.close()


# =====================================================
# GET DATABASE
# =====================================================

def get_db():

    return DatabaseConnection()


# =====================================================
# CREATE DATABASE TABLE
# =====================================================

def init_database():

    conn = get_db()

    try:

        if USE_POSTGRES:

            conn.execute("""
                CREATE TABLE IF NOT EXISTS students (
                    id SERIAL PRIMARY KEY,
                    name TEXT NOT NULL,
                    mobile TEXT NOT NULL,
                    section TEXT NOT NULL,
                    set_number INTEGER NOT NULL,
                    joining_date TEXT NOT NULL,
                    expiry_date TEXT NOT NULL,
                    fees DOUBLE PRECISION NOT NULL,
                    payment_mode TEXT NOT NULL,
                    notes TEXT,
                    created_at TEXT NOT NULL
                )
            """)

        else:

            conn.execute("""
                CREATE TABLE IF NOT EXISTS students (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT NOT NULL,
                    mobile TEXT NOT NULL,
                    section TEXT NOT NULL,
                    set_number INTEGER NOT NULL,
                    joining_date TEXT NOT NULL,
                    expiry_date TEXT NOT NULL,
                    fees REAL NOT NULL,
                    payment_mode TEXT NOT NULL,
                    notes TEXT,
                    created_at TEXT NOT NULL
                )
            """)

        if USE_POSTGRES:

            conn.execute("""
                CREATE TABLE IF NOT EXISTS admins (
                    id SERIAL PRIMARY KEY,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
            """)

        else:

            conn.execute("""
                CREATE TABLE IF NOT EXISTS admins (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
            """)

        conn.commit()


        # -------------------------------------------------
        # SEED 3 DEFAULT ADMIN ACCOUNTS (only if none exist)
        # -------------------------------------------------

        existing_admin = conn.execute(
            "SELECT id FROM admins LIMIT 1"
        ).fetchone()

        if not existing_admin:

            default_admins = [
                ("admin1", "Admin@123"),
                ("admin2", "Admin@123"),
                ("admin3", "Admin@123")
            ]

            created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            for username, password in default_admins:

                conn.execute("""
                    INSERT INTO admins (username, password_hash, created_at)
                    VALUES (?, ?, ?)
                """, (
                    username,
                    generate_password_hash(password),
                    created_at
                ))

            conn.commit()

            print("Seeded 3 default admin accounts: admin1, admin2, admin3 (password: Admin@123)")
            print("IMPORTANT: Change these passwords after first login.")


        print("Database initialized successfully.")

        if USE_POSTGRES:
            print("Database type: PostgreSQL")
        else:
            print("Database type: SQLite")

    except Exception as e:

        print("Database initialization error:", e)

        raise

    finally:

        conn.close()


# =====================================================
# LOGIN REQUIRED DECORATOR
# =====================================================

def login_required(view_function):

    @functools.wraps(view_function)
    def wrapper(*args, **kwargs):

        if not session.get("admin_id"):

            if request.path.startswith("/api/"):

                return jsonify({
                    "success": False,
                    "message": "Please log in first."
                }), 401

            return redirect(url_for("login"))

        return view_function(*args, **kwargs)

    return wrapper


# =====================================================
# LOGIN / LOGOUT
# =====================================================

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "GET":

        return render_template("login.html", error=None)


    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")


    conn = get_db()

    try:

        admin = conn.execute("""
            SELECT id, username, password_hash
            FROM admins
            WHERE username = ?
        """, (username,)).fetchone()

    finally:

        conn.close()


    if admin and check_password_hash(admin["password_hash"], password):

        session["admin_id"] = admin["id"]
        session["admin_username"] = admin["username"]

        return redirect(url_for("index"))


    return render_template(
        "login.html",
        error="Invalid username or password."
    )


@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("login"))


# =====================================================
# ADMIN MANAGEMENT
# =====================================================

@app.route("/api/admins", methods=["GET"])
@login_required
def get_admins():

    conn = get_db()

    try:

        admins = conn.execute("""
            SELECT id, username, created_at
            FROM admins
            ORDER BY id ASC
        """).fetchall()

    finally:

        conn.close()


    result = []

    for admin in admins:

        result.append({
            "id": admin["id"],
            "username": admin["username"],
            "createdAt": admin["created_at"],
            "isYou": admin["id"] == session.get("admin_id")
        })

    return jsonify(result)


@app.route("/api/admins", methods=["POST"])
@login_required
def add_admin():

    data = request.get_json(silent=True)

    if data is None:
        return jsonify({
            "success": False,
            "message": "Invalid or missing JSON data."
        }), 400

    username = str(data.get("username", "") or "").strip()
    password = str(data.get("password", "") or "")


    if not username:
        return jsonify({
            "success": False,
            "message": "Username is required."
        }), 400

    if len(password) < 6:
        return jsonify({
            "success": False,
            "message": "Password must be at least 6 characters."
        }), 400


    conn = get_db()

    try:

        existing = conn.execute("""
            SELECT id FROM admins WHERE username = ?
        """, (username,)).fetchone()

        if existing:
            return jsonify({
                "success": False,
                "message": "This username is already taken."
            }), 409


        created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        conn.execute("""
            INSERT INTO admins (username, password_hash, created_at)
            VALUES (?, ?, ?)
        """, (
            username,
            generate_password_hash(password),
            created_at
        ))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Admin added successfully."
        })

    except Exception as e:

        print("Error adding admin:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while adding the admin."
        }), 500

    finally:

        conn.close()


@app.route("/api/admins/<int:admin_id>", methods=["DELETE"])
@login_required
def delete_admin(admin_id):

    conn = get_db()

    try:

        total = conn.execute("""
            SELECT COUNT(*) AS count FROM admins
        """).fetchone()["count"]

        if total <= 1:
            return jsonify({
                "success": False,
                "message": "Cannot delete the last remaining admin."
            }), 400


        if admin_id == session.get("admin_id"):
            return jsonify({
                "success": False,
                "message": "You cannot delete your own account while logged in."
            }), 400


        admin = conn.execute("""
            SELECT id FROM admins WHERE id = ?
        """, (admin_id,)).fetchone()

        if not admin:
            return jsonify({
                "success": False,
                "message": "Admin not found."
            }), 404


        conn.execute("""
            DELETE FROM admins WHERE id = ?
        """, (admin_id,))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Admin removed successfully."
        })

    except Exception as e:

        print("Error deleting admin:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while removing the admin."
        }), 500

    finally:

        conn.close()


@app.route("/api/change-password", methods=["POST"])
@login_required
def change_password():

    data = request.get_json(silent=True)

    if data is None:
        return jsonify({
            "success": False,
            "message": "Invalid or missing JSON data."
        }), 400

    current_password = str(data.get("currentPassword", "") or "")
    new_password = str(data.get("newPassword", "") or "")


    if len(new_password) < 6:
        return jsonify({
            "success": False,
            "message": "New password must be at least 6 characters."
        }), 400


    conn = get_db()

    try:

        admin = conn.execute("""
            SELECT id, password_hash FROM admins WHERE id = ?
        """, (session.get("admin_id"),)).fetchone()

        if not admin or not check_password_hash(admin["password_hash"], current_password):

            return jsonify({
                "success": False,
                "message": "Current password is incorrect."
            }), 400


        conn.execute("""
            UPDATE admins SET password_hash = ? WHERE id = ?
        """, (
            generate_password_hash(new_password),
            session.get("admin_id")
        ))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Password changed successfully."
        })

    except Exception as e:

        print("Error changing password:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while changing the password."
        }), 500

    finally:

        conn.close()


@app.route("/api/admins/<int:admin_id>/reset-password", methods=["POST"])
@login_required
def reset_admin_password(admin_id):

    data = request.get_json(silent=True)

    if data is None:
        return jsonify({
            "success": False,
            "message": "Invalid or missing JSON data."
        }), 400

    new_password = str(data.get("newPassword", "") or "")


    if len(new_password) < 6:
        return jsonify({
            "success": False,
            "message": "New password must be at least 6 characters."
        }), 400


    conn = get_db()

    try:

        admin = conn.execute("""
            SELECT id FROM admins WHERE id = ?
        """, (admin_id,)).fetchone()

        if not admin:
            return jsonify({
                "success": False,
                "message": "Admin not found."
            }), 404


        conn.execute("""
            UPDATE admins SET password_hash = ? WHERE id = ?
        """, (
            generate_password_hash(new_password),
            admin_id
        ))

        conn.commit()

        return jsonify({
            "success": True,
            "message": "Password reset successfully."
        })

    except Exception as e:

        print("Error resetting password:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while resetting the password."
        }), 500

    finally:

        conn.close()


# =====================================================
# HOME PAGE
# =====================================================

@app.route("/")
@login_required
def index():

    return render_template("index.html")


# =====================================================
# GET ALL STUDENTS
# =====================================================

@app.route("/api/students", methods=["GET"])
@login_required
def get_students():

    today = datetime.now().strftime("%Y-%m-%d")

    conn = get_db()

    try:

        students = conn.execute("""
            SELECT *
            FROM students
            ORDER BY id DESC
        """).fetchall()

        result = []

        for student in students:

            status = (
                "expired"
                if student["expiry_date"] < today
                else "active"
            )

            result.append({
                "id": student["id"],
                "name": student["name"],
                "mobile": student["mobile"],
                "section": student["section"],
                "setNumber": student["set_number"],
                "joiningDate": student["joining_date"],
                "expiryDate": student["expiry_date"],
                "fees": student["fees"],
                "paymentMode": student["payment_mode"],
                "notes": student["notes"] or "",
                "status": status
            })

        return jsonify(result)

    except Exception as e:

        print("Error getting students:", e)

        return jsonify({
            "success": False,
            "message": "Unable to load students."
        }), 500

    finally:

        conn.close()


# =====================================================
# ADD NEW STUDENT
# =====================================================

@app.route("/api/students", methods=["POST"])
@login_required
def add_student():

    data = request.get_json(silent=True)

    if data is None:

        return jsonify({
            "success": False,
            "message": "Invalid or missing JSON data."
        }), 400


    name = str(data.get("name", "") or "").strip()
    mobile = str(data.get("mobile", "") or "").strip()
    section = str(data.get("section", "") or "").strip()
    set_number = data.get("setNumber")
    joining_date = str(
        data.get("joiningDate", "") or ""
    ).strip()

    expiry_date = str(
        data.get("expiryDate", "") or ""
    ).strip()

    fees = data.get("fees")

    payment_mode = str(
        data.get("paymentMode", "") or ""
    ).strip()

    notes = str(
        data.get("notes", "") or ""
    ).strip()


    # -------------------------------------------------
    # BASIC VALIDATION
    # -------------------------------------------------

    if not name:

        return jsonify({
            "success": False,
            "message": "Student name is required."
        }), 400


    if not mobile:

        return jsonify({
            "success": False,
            "message": "Mobile number is required."
        }), 400


    if not mobile.isdigit() or not (7 <= len(mobile) <= 15):

        return jsonify({
            "success": False,
            "message": "Enter a valid mobile number (digits only)."
        }), 400


    if fees is None or str(fees).strip() == "":

        return jsonify({
            "success": False,
            "message": "Fees amount is required."
        }), 400


    try:

        fees = float(fees)

    except (TypeError, ValueError):

        return jsonify({
            "success": False,
            "message": "Fees must be a valid number."
        }), 400


    if fees < 0:

        return jsonify({
            "success": False,
            "message": "Fees cannot be negative."
        }), 400


    if section not in ["General", "VIP", "VVIP"]:

        return jsonify({
            "success": False,
            "message": "Invalid section."
        }), 400


    if not joining_date:

        return jsonify({
            "success": False,
            "message": "Joining date is required."
        }), 400


    if not expiry_date:

        return jsonify({
            "success": False,
            "message": "Expiry date is required."
        }), 400


    if not payment_mode:

        return jsonify({
            "success": False,
            "message": "Payment mode is required."
        }), 400


    if not set_number:

        return jsonify({
            "success": False,
            "message": "Set number is required."
        }), 400


    try:

        set_number = int(set_number)

    except (TypeError, ValueError):

        return jsonify({
            "success": False,
            "message": "Invalid set number."
        }), 400


    # -------------------------------------------------
    # CHECK SET LIMIT
    # -------------------------------------------------

    if section == "General" and not (1 <= set_number <= 106):

        return jsonify({
            "success": False,
            "message": "Invalid General set number."
        }), 400


    if section == "VIP" and not (1 <= set_number <= 57):

        return jsonify({
            "success": False,
            "message": "Invalid VIP set number."
        }), 400


    if section == "VVIP" and not (1 <= set_number <= 12):

        return jsonify({
            "success": False,
            "message": "Invalid VVIP set number."
        }), 400


    # -------------------------------------------------
    # CHECK SET OCCUPIED
    # -------------------------------------------------

    conn = get_db()

    try:

        existing = conn.execute("""
            SELECT id
            FROM students
            WHERE section = ?
            AND set_number = ?
            AND expiry_date >= ?
        """, (
            section,
            set_number,
            datetime.now().strftime("%Y-%m-%d")
        )).fetchone()


        if existing:

            return jsonify({
                "success": False,
                "message": "This set is already occupied."
            }), 409


        # -------------------------------------------------
        # INSERT STUDENT
        # -------------------------------------------------

        created_at = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )


        if USE_POSTGRES:

            cursor = conn.execute("""
                INSERT INTO students (
                    name,
                    mobile,
                    section,
                    set_number,
                    joining_date,
                    expiry_date,
                    fees,
                    payment_mode,
                    notes,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                RETURNING id
            """, (
                name,
                mobile,
                section,
                set_number,
                joining_date,
                expiry_date,
                fees,
                payment_mode,
                notes,
                created_at
            ))

            student_id = cursor.fetchone()["id"]

        else:

            cursor = conn.execute("""
                INSERT INTO students (
                    name,
                    mobile,
                    section,
                    set_number,
                    joining_date,
                    expiry_date,
                    fees,
                    payment_mode,
                    notes,
                    created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                name,
                mobile,
                section,
                set_number,
                joining_date,
                expiry_date,
                fees,
                payment_mode,
                notes,
                created_at
            ))

            student_id = cursor.lastrowid


        conn.commit()


        return jsonify({
            "success": True,
            "message": "Student added successfully.",
            "id": student_id
        })


    except Exception as e:

        print("Error adding student:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while saving the student. Please try again."
        }), 500

    finally:

        conn.close()


# =====================================================
# DELETE / REMOVE STUDENT
# =====================================================

@app.route("/api/students/<int:student_id>", methods=["DELETE"])
@login_required
def delete_student(student_id):

    conn = get_db()

    try:

        student = conn.execute("""
            SELECT id
            FROM students
            WHERE id = ?
        """, (student_id,)).fetchone()


        if not student:

            return jsonify({
                "success": False,
                "message": "Student not found."
            }), 404


        conn.execute("""
            DELETE FROM students
            WHERE id = ?
        """, (student_id,))


        conn.commit()


        return jsonify({
            "success": True,
            "message": "Student removed successfully."
        })


    except Exception as e:

        print("Error removing student:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while removing the student. Please try again."
        }), 500

    finally:

        conn.close()


# =====================================================
# DASHBOARD STATISTICS
# =====================================================

@app.route("/api/dashboard", methods=["GET"])
@login_required
def dashboard():

    today = datetime.now().strftime("%Y-%m-%d")

    conn = get_db()

    try:

        total_students = conn.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE expiry_date >= ?
        """, (today,)).fetchone()[0]


        general_occupied = conn.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE section = 'General'
            AND expiry_date >= ?
        """, (today,)).fetchone()[0]


        vip_occupied = conn.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE section = 'VIP'
            AND expiry_date >= ?
        """, (today,)).fetchone()[0]


        vvip_occupied = conn.execute("""
            SELECT COUNT(*)
            FROM students
            WHERE section = 'VVIP'
            AND expiry_date >= ?
        """, (today,)).fetchone()[0]


        total_sets = 175

        occupied_sets = (
            general_occupied +
            vip_occupied +
            vvip_occupied
        )

        available_sets = (
            total_sets -
            occupied_sets
        )

        general_available = (
            106 -
            general_occupied
        )

        vip_available = (
            57 -
            vip_occupied
        )

        vvip_available = (
            12 -
            vvip_occupied
        )


        return jsonify({

            "totalSets": total_sets,

            "occupiedSets": occupied_sets,

            "availableSets": available_sets,

            "totalStudents": total_students,

            "generalOccupied": general_occupied,

            "generalAvailable": general_available,

            "vipOccupied": vip_occupied,

            "vipAvailable": vip_available,

            "vvipOccupied": vvip_occupied,

            "vvipAvailable": vvip_available

        })

    except Exception as e:

        print("Dashboard error:", e)

        return jsonify({
            "success": False,
            "message": "Unable to load dashboard."
        }), 500

    finally:

        conn.close()


# =====================================================
# EXPORT STUDENTS REPORT
# =====================================================

@app.route("/api/reports/export", methods=["GET"])
@login_required
def export_report():

    from_date = request.args.get(
        "from",
        ""
    ).strip()

    to_date = request.args.get(
        "to",
        ""
    ).strip()

    today = datetime.now().strftime("%Y-%m-%d")


    # -------------------------------------------------
    # VALIDATE DATE RANGE
    # -------------------------------------------------

    if from_date and to_date and from_date > to_date:

        return jsonify({
            "success": False,
            "message": "'From' date cannot be after 'To' date."
        }), 400


    conn = get_db()

    try:

        query = """
            SELECT *
            FROM students
            WHERE 1=1
        """

        params = []


        if from_date:

            query += """
                AND joining_date >= ?
            """

            params.append(from_date)


        if to_date:

            query += """
                AND joining_date <= ?
            """

            params.append(to_date)


        query += """
            ORDER BY joining_date ASC, id ASC
        """


        rows = conn.execute(
            query,
            params
        ).fetchall()


    except Exception as e:

        print("Error generating report:", e)

        return jsonify({
            "success": False,
            "message": "Something went wrong while generating the report."
        }), 500

    finally:

        conn.close()


    # -------------------------------------------------
    # BUILD EXCEL WORKBOOK
    # -------------------------------------------------

    wb = Workbook()

    ws = wb.active

    ws.title = "Students Report"


    headers = [
        "Name",
        "Mobile",
        "Section",
        "Set Number",
        "Joining Date",
        "Expiry Date",
        "Status",
        "Fees",
        "Payment Mode",
        "Notes"
    ]


    header_font = Font(
        name="Arial",
        bold=True,
        color="FFFFFF",
        size=11
    )

    header_fill = PatternFill(
        start_color="2563EB",
        end_color="2563EB",
        fill_type="solid"
    )

    normal_font = Font(
        name="Arial",
        size=11
    )

    bold_font = Font(
        name="Arial",
        size=11,
        bold=True
    )


    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=1,
            column=col,
            value=header
        )

        cell.font = header_font

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    ws.row_dimensions[1].height = 22


    total_fees = 0.0


    for row_idx, student in enumerate(
        rows,
        start=2
    ):

        fees_value = student["fees"] or 0

        total_fees += float(fees_value)


        status_value = (
            "Expired"
            if student["expiry_date"] < today
            else "Active"
        )


        values = [

            student["name"],

            student["mobile"],

            student["section"],

            student["set_number"],

            student["joining_date"],

            student["expiry_date"],

            status_value,

            fees_value,

            student["payment_mode"],

            student["notes"] or ""

        ]


        for col_idx, value in enumerate(
            values,
            start=1
        ):

            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=value
            )

            cell.font = normal_font


    column_widths = [
        22,
        15,
        10,
        12,
        14,
        14,
        10,
        10,
        14,
        28
    ]


    for i, width in enumerate(
        column_widths,
        start=1
    ):

        ws.column_dimensions[
            get_column_letter(i)
        ].width = width


    # -------------------------------------------------
    # SUMMARY
    # -------------------------------------------------

    summary_start = len(rows) + 3

    range_label = "All time"


    if from_date or to_date:

        range_label = (
            f"{from_date or 'Start'} "
            f"to "
            f"{to_date or 'Today'}"
        )


    ws.cell(
        row=summary_start,
        column=1,
        value="Date Range:"
    ).font = bold_font

    ws.cell(
        row=summary_start,
        column=2,
        value=range_label
    ).font = normal_font


    ws.cell(
        row=summary_start + 1,
        column=1,
        value="Total Students:"
    ).font = bold_font

    ws.cell(
        row=summary_start + 1,
        column=2,
        value=len(rows)
    ).font = normal_font


    ws.cell(
        row=summary_start + 2,
        column=1,
        value="Total Fees Collected:"
    ).font = bold_font

    ws.cell(
        row=summary_start + 2,
        column=2,
        value=round(total_fees, 2)
    ).font = normal_font


    # -------------------------------------------------
    # SEND EXCEL FILE
    # -------------------------------------------------

    output = BytesIO()

    wb.save(output)

    output.seek(0)


    filename_range = (
        f"{from_date or 'all'}"
        f"_to_"
        f"{to_date or 'all'}"
    )

    filename = (
        f"students_report_"
        f"{filename_range}.xlsx"
    )


    return send_file(

        output,

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),

        as_attachment=True,

        download_name=filename

    )


# =====================================================
# INITIALIZE DATABASE
# =====================================================

init_database()


# =====================================================
# RUN APPLICATION
# =====================================================

if __name__ == "__main__":

    print("")

    print("======================================")

    print("      STUDY POINT MANAGEMENT SYSTEM")

    print("======================================")

    print("")

    print("Local Server Started")

    print("Laptop: http://127.0.0.1:5000")

    print("")

    if USE_POSTGRES:

        print("Database: PostgreSQL")

    else:

        print("Database: SQLite")


    print("")


    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )
